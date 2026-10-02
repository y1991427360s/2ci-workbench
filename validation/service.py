"""Read-only adapters for existing project files and atomic validation reports."""
from __future__ import annotations

import hashlib
import json
import zipfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from cable_stat.loaders import find_header_row, _merged_value_lookup, load_aliases, load_cabinets, MODEL_HEADERS, SPEC_HEADERS
from cable_stat.project import scan_project
from cable_stat.csvio import read_csv_dicts
from cable_stat.text import normalize_cabinet_name, normalize_text
from modules.terminal import duanzi_dxf_tool as drawing
from modules.terminal.service import TerminalService, _atomic_write_files
from .engine import validate


def input_files(path):
    files = sorted(p for p in path.rglob('*') if p.is_file() and
                  'outputs' not in p.relative_to(path).parts and p.suffix.lower() in {'.xlsx', '.csv', '.txt', '.json'}
                  and p.name != 'result.json')
    for p in files:
        if not p.resolve().is_relative_to(path.resolve()):
            raise ValueError('工程输入路径越界：' + p.relative_to(path).as_posix())
    return files


def fingerprint(path):
    digest = hashlib.sha256()
    for p in input_files(path):
        digest.update(p.relative_to(path).as_posix().encode())
        digest.update(p.read_bytes())
    result = path / 'result.json'
    if result.exists():
        if not result.resolve().is_relative_to(path.resolve()):
            raise ValueError('路径结果文件越界')
        digest.update(result.read_bytes())
    return digest.hexdigest()


def record(raw, file, row, kind, **fields):
    defaults = dict(cable_number='', source_cabinet='', target_cabinet='', terminal_number='', terminal_strip='')
    return {**defaults, **fields, 'entity_type': kind, 'entity_id': f'{file}:{row}',
            'file': file, 'row': row, 'raw': raw}


def csv_rows(path):
    return read_csv_dicts(path)


def load_cables(path):
    candidates = []
    for p in sorted(path.glob('*.xlsx')):
        if p.name.startswith('~$'):
            continue
        wb = openpyxl.load_workbook(p, data_only=False)
        try:
            ws = wb.active
            header = find_header_row(list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True)),
                                     ('电缆编号', '起点', '终点'))
            if header:
                candidates.append(p)
        finally:
            wb.close()
    csv_path = path / 'data' / '电缆清册.csv'
    if csv_path.exists():
        candidates.append(csv_path)
    if len(candidates) != 1:
        raise ValueError('需要唯一电缆清册（XLSX 或 电缆清册.csv），当前候选：' + '、'.join(p.name for p in candidates))
    p = candidates[0]
    filename = p.relative_to(path).as_posix()
    rows = []
    if p.suffix == '.csv':
        raw_rows = csv_rows(p)
        if not raw_rows or not {'电缆编号', '起点', '终点'}.issubset(raw_rows[0]):
            raise ValueError('电缆清册.csv 缺少表头：电缆编号、起点、终点')
        entries = enumerate(raw_rows, 2)
        for n, raw in entries:
            rows.append(record(raw, filename, n, 'cable', cable_number=normalize_text(raw['电缆编号']),
                               source_cabinet=normalize_cabinet_name(raw['起点']), target_cabinet=normalize_cabinet_name(raw['终点']),
                               specification=normalize_text(raw.get('规格') or raw.get('电缆规格') or raw.get('型号'))))
    else:
        wb = openpyxl.load_workbook(p, data_only=False)
        try:
            ws = wb.active
            n, headers = find_header_row(list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True)), ('电缆编号', '起点', '终点'))
            merged = _merged_value_lookup(ws, set(headers.values()))
            for index in range(n + 1, ws.max_row + 1):
                raw = {key: merged.get((index, col), ws.cell(index, col).value) for key, col in headers.items()}
                if not any(raw.get(k) is not None for k in ('电缆编号', '起点', '终点')):
                    continue
                if normalize_text(raw.get('电缆编号')) == '电缆编号':
                    continue
                # Full-width merged section headings are not cable declarations.
                if any(a.min_col < a.max_col and a.min_row <= index <= a.max_row and
                       sum(a.min_col <= headers[k] <= a.max_col for k in ('电缆编号', '起点', '终点')) >= 2
                       for a in ws.merged_cells.ranges):
                    continue
                spec = ' '.join(normalize_text(raw[k]) for k in dict.fromkeys((*MODEL_HEADERS, *SPEC_HEADERS, '电缆规格', '电缆型号规格')) if raw.get(k))
                rows.append(record(raw, filename, index, 'cable', cable_number=normalize_text(raw.get('电缆编号')),
                                   source_cabinet=normalize_cabinet_name(raw.get('起点')), target_cabinet=normalize_cabinet_name(raw.get('终点')), specification=spec))
        finally:
            wb.close()
    if not rows:
        raise ValueError('电缆清册没有有效数据行')
    return rows


def load_terminals(path, ctx):
    structured = path / 'data' / 'terminal' / '端子数据.csv'
    if structured.exists():
        raw_rows = csv_rows(structured)
        required = {'terminal_strip', 'terminal_number', 'cable_number', 'source_cabinet', 'target_cabinet'}
        if not raw_rows or not required.issubset(raw_rows[0]):
            raise ValueError('端子数据.csv 必需表头：' + ', '.join(sorted(required)))
        rows = []
        for n, raw in enumerate(raw_rows, 2):
            fields = {k: normalize_text(v) for k, v in raw.items() if k}
            for k in ('source_cabinet', 'target_cabinet'):
                fields[k] = normalize_cabinet_name(fields.get(k))
            fields['external'] = fields.get('external', '').lower() in {'true', '1', 'yes', '是'} or bool(fields.get('target_cabinet') or fields.get('cable_number'))
            fields['valid_connection'] = fields.get('valid_connection', '').lower() in {'true', '1', 'yes', '是'}
            rows.append(record(raw, 'data/terminal/端子数据.csv', n, 'terminal', **fields))
        ctx['coverage']['terminal_source'] = '端子数据.csv（优先于出图 TXT；每行一个物理端子，source_cabinet 为本柜侧）'
        return rows
    service = TerminalService(path)
    state = service.load_state()
    for warning in service.load_warnings:
        ctx['input_issues'].append(dict(message=warning, entity_id='data/terminal'))
    ctx['coverage']['terminal_source'] = '原端子 TXT（单柜侧，电缆编号由出图规则自动生成）'
    if not state['terminals'].strip():
        ctx['input_issues'].append(dict(message='未提供端子数据，端子覆盖检查不完整', severity='WARNING', entity_id='data/terminal'))
        return []
    blocks, errors = drawing.parse_terminals(state['terminals'], preserve_duplicates=True)
    connections, _, wire_errors, warnings = drawing.parse_wiring(state['wiring'], blocks, state['cabinet'], drawing.DEFAULTS['cablePrefix'])
    for message in errors + wire_errors:
        ctx['input_issues'].append(dict(message=message, entity_id='data/terminal', raw=state))
    for message in warnings:
        ctx['input_issues'].append(dict(message=message, severity='WARNING', entity_id='data/terminal/接线.txt', raw=state['wiring']))
    rows = []
    for b, block in enumerate(blocks, 1):
        for n, number in enumerate(block['terminals'], 1):
            key = (block['name'], number)
            matched = [c for c in connections if (c['terminalBlock'], c['terminal']) == key] or [{}]
            for c in matched:
                rows.append(record({'block': block, 'connection': c, 'text': state['terminals'], 'wiring': state['wiring']},
                                   'data/terminal/端子排.txt', f'{b}:{n}', 'terminal', declaration_id=f'{b}:{n}',
                                   terminal_strip=block['name'], terminal_number=number,
                                   source_cabinet=normalize_cabinet_name(state['cabinet']), target_cabinet=normalize_cabinet_name(c.get('toCabinet')),
                                   cable_number=c.get('cableNumber', ''), circuit_number=c.get('principle', ''), external=bool(c)))
    return rows


def load_project(path):
    path = Path(path)
    ctx = dict(project_id=path.name, cables=[], terminals=[], input_issues=[], mapping_issues=[], path_issues=[], coverage={})
    for key, loader in (('cables', lambda: load_cables(path)), ('terminals', lambda: load_terminals(path, ctx))):
        try:
            ctx[key] = loader()
        except (ValueError, OSError, KeyError, TypeError, zipfile.BadZipFile) as error:
            ctx['input_issues'].append(dict(message=f'{key} 输入无法读取：{error}', entity_id=key))
    try:
        alias_map, warnings = load_aliases(path / 'data')
        cabinets, cabinet_warnings = load_cabinets(path / 'data')
    except (ValueError, OSError) as error:
        alias_map, cabinets, warnings, cabinet_warnings = {}, {}, [], []
        ctx['input_issues'].append(dict(message=f'映射输入无法读取：{error}', entity_id='data/柜名别名.csv'))
    for warning in warnings + cabinet_warnings:
        ctx['mapping_issues'].append(dict(message=warning, entity_id='data/柜名别名.csv',
                                         raw_aliases=csv_rows(path / 'data' / '柜名别名.csv'),
                                         raw_cabinets=csv_rows(path / 'data' / '柜子坐标.csv')))
    for source, target in alias_map.items():
        if target not in cabinets:
            ctx['mapping_issues'].append(dict(message=f'映射目标不存在：{source} → {target}', entity_id='data/柜名别名.csv', raw={'清册名称': source, 'CAD名称': target}))
    for row in ctx['cables'] + ctx['terminals']:
        for key in ('source_cabinet', 'target_cabinet'):
            row[key] = alias_map.get(row[key], row[key]).casefold()
    status = scan_project(path)
    if not (path / 'result.json').exists() or status.result_outdated:
        ctx['path_issues'].append(dict(message='尚无路径计算结果或结果已过期，请重新计算；本次未验证路径正确性', severity='INFO', entity_id='result.json'))
    else:
        try:
            result = json.loads((path / 'result.json').read_text(encoding='utf-8'))
            for n, issue in enumerate(result.get('issues', []), 1):
                ctx['path_issues'].append(dict(message=str(issue), raw=issue, entity_id=f'result.json:issues:{n}'))
            if result.get('failed'):
                ctx['path_issues'].append(dict(message=f"现有计算结果有 {result['failed']} 根电缆未计算成功", raw=result, entity_id='result.json'))
        except (ValueError, OSError, AttributeError) as error:
            ctx['input_issues'].append(dict(message=f'路径结果无法读取：{error}', entity_id='result.json'))
    ctx['coverage'].update(cable_rows=len(ctx['cables']), terminal_rows=len(ctx['terminals']),
                           cabinet_coordinates=len(cabinets), cabinet_aliases=len(alias_map),
                           path_result_reused=(path / 'result.json').exists() and not status.result_outdated)
    return ctx


def markdown(report):
    lines = ['# 二次设计项目一致性检查 V2', '', f"工程：{report['project_name']}（{report['project_id']}）", f"检查时间：{report['generated_at']}",
             '', ' / '.join(f'{k}: {v}' for k, v in report['counts'].items()), '', '## 数据覆盖', '',
             json.dumps(report['coverage'], ensure_ascii=False), '', '提示：0 ERROR 表示已加载数据未发现确定性错误，仍需查看覆盖范围及 WARNING / INFO。', '']
    for category, count in report['categories'].items():
        lines.extend([f'## {category}（{count}）', ''])
        for issue in report['issues']:
            if issue['category'] != category:
                continue
            lines.extend([f"### [{issue['severity']}] {issue['rule_id']} {issue['message']}", '',
                          f"对象：{issue['entity_id']}；电缆：{issue['cable_number'] or '—'}；端子：{issue['terminal_number'] or '—'}",
                          f"起终点：{issue['source'] or '—'} → {issue['target'] or '—'}", '', '原始证据：', '```json',
                          json.dumps(issue['evidence'], ensure_ascii=False, indent=2, default=str), '```', '', '处理建议：' + issue['suggestion'], ''])
    return '\n'.join(lines)


def run_project(path):
    path = Path(path)
    signature = fingerprint(path)
    directory = path / 'outputs' / 'reports'
    if not directory.resolve().is_relative_to(path.resolve()):
        raise ValueError('报告输出路径越界')
    for name in ('validation.json', 'validation.md'):
        if not (directory / name).resolve().is_relative_to(path.resolve()):
            raise ValueError('报告文件路径越界')
    report = validate(load_project(path))
    report.update(version=2, project_name=json.loads((path / 'project.json').read_text(encoding='utf-8'))['name'],
                  generated_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds'), input_fingerprint=signature, stale=False)
    _atomic_write_files({directory / 'validation.json': json.dumps(report, ensure_ascii=False, indent=2, default=str),
                         directory / 'validation.md': markdown(report)})
    return report


def read_report(path):
    path = Path(path)
    file = path / 'outputs' / 'reports' / 'validation.json'
    if not file.resolve().is_relative_to(path.resolve()):
        raise ValueError('报告文件路径越界')
    if not file.exists():
        return None
    report = json.loads(file.read_text(encoding='utf-8'))
    report['stale'] = report.get('input_fingerprint') != fingerprint(path)
    return report
